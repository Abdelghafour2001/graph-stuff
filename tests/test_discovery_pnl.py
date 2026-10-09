"""Structure discovery, P&L engine, market references, concept matching on a made-up branch workbook. No Neo4j, no LLM."""
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).parent.parent
sys.path[:0] = [str(ROOT / "src"), str(ROOT / "scripts")]
import concept_match  # noqa: E402
import discovery  # noqa: E402
import gateway  # noqa: E402
import market  # noqa: E402
import pnl  # noqa: E402
import sample_branch  # noqa: E402


@pytest.fixture(scope="module")
def found(tmp_path_factory):
    path = tmp_path_factory.mktemp("branch") / "branch.xlsx"
    sample_branch.build(path)
    return discovery.discover(path)


def rows(found, sheet):
    return {r["label"]: r for r in found["sheets"][sheet]["rows"]}


def test_formulas_inputs_and_meaning(found):
    r = rows(found, "Budget S1 2025")
    assert r["Marge brute (M$)"]["role"] == "computed" and "[Coûts fixes (M$)]" in r["Marge brute (M$)"]["formula"]
    assert r["Production DAP (kt)"]["role"] == "input" and r["Production DAP (kt)"]["driver"] == "volume"
    assert r["Consommation spécifique soufre (t/t DAP)"]["concept"]["concept"] == "sulfur"
    assert r["Prix soufre CFR Jorf ($/t)"]["concept"]["concept"] == "sulfur" and r["Prix soufre CFR Jorf ($/t)"]["site"] == "jorf_lasfar"
    assert r["Coût soufre (M$)"]["overrides"] == ["E"]


def test_pasted_values_logic_is_found_by_arithmetic(found):
    ids = {i["target"]: i for i in found["sheets"]["Synthèse (valeurs)"]["identities"] if i["kind"] != "flat"}
    assert ids["Marge brute (M$)"]["kind"] == "sum" and len(ids["Marge brute (M$)"]["terms"]) == 4
    assert ids["Chiffre d'affaires (M$)"]["kind"] == "product" and ids["Chiffre d'affaires (M$)"]["k"] == pytest.approx(0.001)
    assert ids["Coût soufre (M$)"]["kind"] == "ratio" and ids["Coût soufre (M$)"]["k"] == pytest.approx(0.066)  # 0.40 t/t x 165 $/t
    assert "Coût ammoniac (M$)" not in ids  # no ammonia price on that sheet: it stays an input, not a guess


def test_questions(found):
    q = " | ".join(found["questions"])
    assert "typed over the formula" in q and "hides the numbers 1.6, 48" in q and "[1]Hypotheses" in q
    assert "phosphoric_acid" in q  # the graph expects DAP's intermediates


def model(found):
    return pnl.from_discovery(found["sheets"]["Budget S1 2025"])


def test_engine_matches_the_branch_formulas(found):
    m = model(found)
    jan = 520 * 650 / 1000 - 520 * 0.4 * 165 / 1000 - 520 * 0.22 * 430 / 1000 - 520 * 1.6 * 48 / 1000 - 42
    assert m.evaluate()["Marge brute (M$)"][0] == pytest.approx(jan)


def test_shapley_bridge_adds_up_and_sensitivities_have_the_right_sign(found):
    m = model(found)
    target = "Marge brute (M$)"
    after = pnl.mark_to_market(m, {"Prix de vente DAP FOB ($/t)": [700] * 6, "Prix ammoniac CFR ($/t)": [500] * 6})
    b = pnl.shapley(m, target, m.inputs, after)
    assert b["total_change"] == pytest.approx(m.total(target, after) - m.total(target))
    assert sum(b["by_input"].values()) == pytest.approx(b["total_change"]) and abs(b["unexplained"]) < 1e-6
    assert b["by_input"]["Prix de vente DAP FOB ($/t)"] > 0 > b["by_input"]["Prix ammoniac CFR ($/t)"]
    sens = {s["input"]: s["per_1pct"] for s in pnl.sensitivities(m, target)}
    assert sens["Prix de vente DAP FOB ($/t)"] > 0 > sens["Prix soufre CFR Jorf ($/t)"]


def test_market_reference_uses_actual_months_only(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Price Forecast"
    ws.append(["$/t", "DAP Morocco fob", None])
    ws.append([None, "Low", "High"])
    for m, fmt in ((9, "mmm-yy"), (10, 'mmm\\-yy\\ "f"')):
        ws.append([datetime(2025, m, 1), 700, 720])
        ws.cell(ws.max_row, 1).number_format = fmt
    wb.save(tmp_path / "Argus Monthly Phosphates Outlook - test.xlsx")
    series, _ = market.reference("dap", "price", tmp_path)
    assert series == {"2025-09": 710}


def test_concept_match_and_thinking_strip(monkeypatch):
    monkeypatch.delenv("BIFROST_RERANK_MODEL", raising=False)
    concepts = concept_match.concepts_from_yaml()
    m = concept_match.match("Prix soufre CFR Jorf ($/t)", concepts, ("input", "product"))
    assert m["concept"] == "sulfur" and m["ranked_by"] == "lexical"
    assert concept_match.match("Taux de change USD/MAD", concepts, ("input", "product"))["decision"] == "unknown"
    assert gateway.strip_thinking("<think>17*23 = 391</think>\n391") == "391"
