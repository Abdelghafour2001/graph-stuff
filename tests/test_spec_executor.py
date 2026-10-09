"""Spec executor on small synthetic workbooks: anchors, forecast status, missing markers, period order. No Neo4j needed."""
import os
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT / "src"))
for k, val in {"NEO4J_URI": "bolt://localhost:7687", "NEO4J_USER": "neo4j", "NEO4J_PASSWORD": "x", "EXCEL_DIR": "."}.items():
    os.environ.setdefault(k, val)
import spec_executor as se  # noqa: E402

FORECAST_FMT = 'mmm\\-yy\\ "f"'


def workbook(path: Path, blank_rows_above: int, typo: bool = False) -> None:
    """A price sheet like Argus: unit cell + series names, then monthly rows; the last two months are forecasts."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Price Forecast"
    top = 2 + blank_rows_above
    ws.cell(1, 1, "Copyright")
    ws.cell(top, 1, "$/t")
    ws.cell(top, 2, "DAP Morocco fob")
    ws.cell(top, 3, "DAP India cfr")
    months = [datetime(2025, m, 1) for m in range(6, 12)]
    if typo:
        months[2] = datetime(2024, 8, 1)  # Aug-24 between Jul-25 and Sep-25
    for i, d in enumerate(months):
        r = top + 1 + i
        ws.cell(r, 1, d).number_format = FORECAST_FMT if d.month >= 10 else "mmm-yy"
        ws.cell(r, 2, 600 + i)
        ws.cell(r, 3, "NM" if i == 1 else 650 + i)
    wb.save(path)


def spec(**extra) -> dict:
    base = {"file": "prices.xlsx", "sheet": "Price Forecast", "table_range": "A2:C8", "header_rows": [2], "label_columns": ["A"],
            "row_dimensions": ["date"], "column_dimension": "route", "drop_rows": [], "drop_columns": [], "skip_rows_matching": [],
            "skip_label_kinds": [], "operators": [{"op": "stack", "args": {}}], "value_unit": "$/t", "fixed": {"product": "dap"},
            "notes": "", "anchor": {"cell": "A2", "text": "$/t"}}
    return {**base, **extra}


@pytest.fixture
def excel(tmp_path, monkeypatch):
    monkeypatch.setattr(se, "EXCEL_DIR", tmp_path)
    monkeypatch.setattr(se, "label_kinds", lambda labels: {})
    return tmp_path


def run(s):
    se.validate(s)
    rows, ctx = se.execute(s)
    return rows, ctx, {c["check"]: c for c in se.check(s, rows, ctx)}


def test_forecast_rows_come_from_the_cell_format(excel):
    workbook(excel / "prices.xlsx", 0)
    rows, ctx, checks = run(spec())
    forecast = {r["date"][:7] for r in rows if r["status"] == "forecast"}
    assert forecast == {"2025-10", "2025-11"}
    assert all(c["ok"] for c in checks.values()), checks


def test_published_date_flags_unmarked_forecasts(excel):
    workbook(excel / "prices.xlsx", 0)
    rows, _, _ = run(spec(published="2025-08-31"))
    assert {r["date"][:7] for r in rows if r["status"] == "forecast"} == {"2025-09", "2025-10", "2025-11"}


def test_anchor_follows_a_shifted_layout_and_refuses_a_changed_one(excel):
    workbook(excel / "prices.xlsx", 1)  # one extra row above the table, as in a later edition
    rows, ctx, checks = run(spec())
    assert ctx["anchor_offset"] == 1 and checks["header_is_text"]["ok"]
    assert rows[0]["route"] == "DAP Morocco fob" and rows[0]["date"].startswith("2025-06")
    with pytest.raises(AssertionError, match="layout changed"):
        se.execute(spec(anchor={"cell": "A2", "text": "kt"}))


def test_missing_markers_are_skipped_not_counted_as_bad_values(excel):
    workbook(excel / "prices.xlsx", 0)
    rows, ctx, checks = run(spec(missing_values=["NM", 0]))
    assert ctx["missing"] == 1 and ctx["non_numeric"] == 0 and checks["values_parse"]["ok"]
    assert len(rows) == 11  # 6 months x 2 routes, minus the NM cell


def test_backwards_period_is_flagged(excel):
    workbook(excel / "prices.xlsx", 0, typo=True)
    _, _, checks = run(spec())
    assert not checks["row_periods_in_order"]["ok"] and "2024-08-01" in checks["row_periods_in_order"]["detail"]


def test_parse_period():
    assert se.parse_period("2Q26").isoformat() == "2026-04-01"
    assert se.parse_period("Q3 2025").isoformat() == "2025-07-01"
    assert se.parse_period("Oct-25").isoformat() == "2025-10-01"
    assert se.parse_period(2027).isoformat() == "2027-01-01"
    assert se.parse_period("Export total") is None
