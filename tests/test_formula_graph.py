"""Formula harvesting on a small synthetic branch workbook. No Neo4j needed."""
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))
import formula_graph as fg  # noqa: E402


def branch_workbook(path: Path) -> None:
    """A branch model with no template: volumes and prices per product, revenue computed, totals, one external input."""
    wb = openpyxl.Workbook()
    vol = wb.active
    vol.title = "Volumes"
    price = wb.create_sheet("Prices")
    rev = wb.create_sheet("Revenue")
    for ws in (vol, price, rev):
        ws.append(["Product", "Jan", "Feb", "Mar"])
    for i, p in enumerate(["DAP", "MAP", "TSP"], start=2):
        vol.append([p, 100 * i, 110 * i, 120 * i])
        price.append([p, f"=[1]Market!B{i + 10}", f"=[1]Market!C{i + 10}", f"=[1]Market!D{i + 10}"])
        rev.append([p] + [f"=Volumes!{c}{i}*Prices!{c}{i}" for c in "BCD"])
    rev.append(["Total"] + [f"=SUM({c}2:{c}4)" for c in "BCD"])
    wb.save(path)


def test_rules_dependencies_and_external_links(tmp_path):
    branch_workbook(tmp_path / "branch.xlsx")
    h = fg.harvest(tmp_path / "branch.xlsx")
    revenue = {r["rule"]: r["rows"] for r in h["sheets"]["Revenue"]["rules"]}
    assert revenue["Volumes[row]*Prices[row]"] == 3          # 9 cell formulas, one rule
    assert revenue["SUM([block above])"] == 1
    assert {r["rule"] for r in h["sheets"]["Prices"]["rules"]} == {"[1]Market[row+10]"}
    assert {(d["sheet"], d["uses"]) for d in h["depends_on"]} == {("Revenue", "Volumes"), ("Revenue", "Prices")}
    assert h["external_links"] == {"[1]Market": 9}


def test_label_column_is_the_most_distinct():
    wb = openpyxl.Workbook()
    ws = wb.active
    for name, status in [("Linz", "On-stream"), ("Antwerp", "On-stream"), ("Tertre", "On-stream")]:
        ws.append([name, status, 1])
    _, labels = fg.sheet_cells(ws)
    assert labels == {1: "Linz", 2: "Antwerp", 3: "Tertre"}
