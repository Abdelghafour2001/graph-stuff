"""Recover the calculation structure of a workbook from its formulas, without a template.

Branches send their own spreadsheets; the logic (volume x price, sums of sites, ratios) is in the formulas. This module
reads every formula, replaces cell references by the row labels they point to, and generalises over columns and rows:
100,000 cell formulas become a handful of readable rules per sheet, e.g.

  Op. rate[row] = IFERROR(Production[row] / 'Capacity by geography'[row], ...)      (300 rows x 40 years)
  Demand[row]   = 'NH3 for urea'[row] + 'NH3 for AN'[row] + ...
  Global        = SUM(rows Austria .. Zimbabwe)

plus the sheet dependency graph and the links to workbooks we do not have ([2]Model). Deterministic; the LLM only names
the rules (maps labels to concepts) and asks the branch about what it cannot see.

  python formula_graph.py <workbook.xlsx> [sheet ...]
"""
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils import column_index_from_string, range_boundaries

LABEL_COLS = 6         # row labels are looked for in the first columns
REF = re.compile(r"^(?:(\[\d+\])?(?:'((?:[^']|'')+)'|([^'!:]+))!)?(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)$")


def sheet_cells(ws, numbers: dict | None = None) -> tuple[dict, dict]:
    """Formulas {(row, col): text} and row labels {row: label} of one sheet.

    The label column is the one, among the first LABEL_COLS, with the most distinct texts: countries rather than regions,
    plants rather than their status."""
    formulas, texts = {}, defaultdict(dict)
    for r, row in enumerate(ws.iter_rows(), start=1):
        for c, cell in enumerate(row, start=1):
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                formulas[(r, c)] = v
            elif isinstance(v, str) and v.strip() and c <= LABEL_COLS:
                texts[c][r] = v.strip()
            elif numbers is not None and isinstance(v, (int, float)) and not isinstance(v, bool):
                numbers[(r, c)] = v
    if not texts:
        return formulas, {}
    best = max(texts, key=lambda c: (len(set(texts[c].values())), -c))
    labels = dict(texts[best])
    for c in sorted(texts):  # rows without a text in the label column (totals, notes) keep their own text
        for r, t in texts[c].items():
            labels.setdefault(r, t)
    return formulas, labels


def references(formula: str) -> list[tuple[str | None, str | None, str]]:
    """(external workbook, sheet, A1 ref or range) for every reference in a formula."""
    out = []
    for tok in Tokenizer(formula).items:
        if tok.type == "OPERAND" and tok.subtype == "RANGE":
            m = REF.match(tok.value)
            if m:
                out.append((m.group(1), (m.group(2) or m.group(3) or "").replace("''", "'") or None, m.group(4).replace("$", "")))
    return out


def abstract(formula: str, sheet: str, row: int, labels: dict) -> tuple[str, list]:
    """The formula with each reference replaced by Sheet[row label]; the same row as the formula becomes [row]."""
    deps = []
    parts = []
    for tok in Tokenizer(formula).items:
        if tok.type == "OPERAND" and tok.subtype == "RANGE" and (m := REF.match(tok.value)):
            ext, ref_sheet, a1 = m.group(1), (m.group(2) or m.group(3) or "").replace("''", "'") or sheet, m.group(4).replace("$", "")
            c1, r1, c2, r2 = range_boundaries(a1)
            name = f"{ext}{ref_sheet}" if ext else ref_sheet
            lab = labels.get(ref_sheet, {})
            prefix = name if ref_sheet != sheet or ext else ""
            if r1 == r2:
                if r1 == row:
                    where = "row"
                elif ext or r1 not in lab:  # unlabelled or external target: keep the offset, it generalises over rows
                    where = f"row{r1 - row:+d}"
                else:
                    where = lab[r1]
                text = f"{prefix}[{where}]"
            elif ref_sheet == sheet and not ext and r2 == row - 1:
                text = "[block above]"
            elif ref_sheet == sheet and not ext and r1 == row + 1:
                text = "[block below]"
            else:
                text = f"{prefix}[{lab.get(r1, f'r{r1}')} .. {lab.get(r2, f'r{r2}')}]"
            deps.append((ext, ref_sheet))
            parts.append(text)
        else:
            parts.append(tok.value if tok.type != "OPERAND" or tok.subtype != "TEXT" else tok.value)
    return "".join(parts), deps


def harvest(path: Path, only: list[str] | None = None) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True)
    names = [n for n in wb.sheetnames if not only or n in only]
    cells, labels, consts = {}, {}, {}
    for n in wb.sheetnames:  # labels of every sheet, so cross-sheet references can be named
        consts[n] = {}
        f, lab = sheet_cells(wb[n], consts[n] if n in names else None)
        labels[n] = lab
        if n in names:
            cells[n] = f
    wb.close()

    rules, edges, external = {}, Counter(), Counter()
    for n, formulas in cells.items():
        per_row = defaultdict(Counter)
        for (r, c), f in formulas.items():
            text, deps = abstract(f, n, r, labels)
            per_row[r][text] += 1
            for ext, s in deps:
                if ext:
                    external[f"{ext}{s}"] += 1
                elif s != n:
                    edges[(n, s)] += 1
        # generalise over rows: identical row-relative rules collapse into one
        sheet_rules = Counter()
        examples = {}
        for r, texts in per_row.items():
            text, count = texts.most_common(1)[0]
            generic = any(t in text for t in ("[row]", "[row+", "[row-", "[block above]", "[block below]"))
            key = text if generic else f"{labels[n].get(r, f'r{r}')} = {text}"
            sheet_rules[key] += 1
            examples.setdefault(key, labels[n].get(r, f"r{r}"))
        # per row: its own expression (for the P&L engine) and the hard-typed numbers inside its formula span (overrides)
        row_rules, by_row_cols = [], defaultdict(list)
        for (r, c) in formulas:
            by_row_cols[r].append(c)
        for r, texts in sorted(per_row.items()):
            expr, count = texts.most_common(1)[0]
            cols = by_row_cols[r]
            overrides = sorted(c for (rr, c) in consts[n] if rr == r and min(cols) < c < max(cols))
            row_rules.append({"row": r, "label": labels[n].get(r, f"r{r}"), "expr": expr, "columns": count,
                              "variants": len(texts), "overrides": [openpyxl.utils.get_column_letter(c) for c in overrides]})
        rules[n] = {"row_rules": row_rules, "formulas": len(formulas), "rules": [{"rule": k if not k.startswith("=") else f"[row] {k}", "rows": v, "example_row": examples[k]}
                                                          for k, v in sheet_rules.most_common()]}
    return {"file": path.name, "sheets": rules,
            "depends_on": [{"sheet": a, "uses": b, "references": n} for (a, b), n in edges.most_common()],
            "external_links": dict(external)}


def summary(h: dict, top: int = 4) -> str:
    lines = [f"{h['file']}"]
    for n, s in h["sheets"].items():
        if not s["formulas"]:
            continue
        lines.append(f"\n  {n}: {s['formulas']} formulas -> {len(s['rules'])} rules")
        for r in s["rules"][:top]:
            lines.append(f"    {r['rows']:>4} rows  {r['rule'][:150]}")
    if h["external_links"]:
        lines.append("\n  links to workbooks we do not have: " + ", ".join(f"{k} ({v})" for k, v in h["external_links"].items()))
    return "\n".join(lines)


if __name__ == "__main__":
    print(summary(harvest(Path(sys.argv[1]), sys.argv[2:] or None)))
