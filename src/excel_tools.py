"""Agent tools over the workbook graph and the raw Excel files. Read-only, except propose_extraction_spec (review queue)."""
import json
import os
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import yaml
from anthropic import beta_tool
from openpyxl.utils import get_column_letter, range_boundaries

import formula_graph
from graph import read_graph
from spec_executor import check, execute, validate

EXCEL_DIR = Path(os.environ["EXCEL_DIR"])
SPECS = Path(__file__).parent.parent / "specs"
MAX_CELLS = 600


def cell_text(v) -> str:
    return "" if v is None else str(v)[:40]


@beta_tool
def find_sheets(concept_ids: list[str]) -> str:
    """Find workbook sheets whose headers, row labels or context mention ALL the given concepts. Data sheets first.

    Args:
        concept_ids: Concept ids from lookup_term, e.g. ["dap", "country_morocco", "fob"].
    """
    rows = read_graph(
        "MATCH (x:Sheet)-[m:MENTIONS]->(c:Concept) WHERE c.id IN $ids "
        "WITH x, collect(DISTINCT c.id) AS hit, sum(m.n) AS strength WHERE size(hit) = size($ids) "
        "RETURN x.file AS file, x.name AS sheet, x.role AS role, x.description AS description, strength "
        "ORDER BY x.role = 'data' DESC, strength DESC",
        ids=concept_ids,
    )
    return json.dumps(rows, ensure_ascii=False) if rows else "No sheet mentions all of these concepts. Try fewer concepts."


@beta_tool
def describe_sheet(file: str, sheet: str) -> str:
    """Layout features, linked concepts, a compact skeleton (first non-empty rows) of one sheet, and hints from the old
    ingestion pipeline (which hand-written caster was meant to read this sheet, and its documentation).

    Args:
        file: Workbook file name as returned by find_sheets.
        sheet: Sheet name.
    """
    rows = read_graph(
        "MATCH (x:Sheet {key: $key}) OPTIONAL MATCH (x)-[m:MENTIONS]->(c:Concept) "
        "WITH x, collect({concept: c.id, where: m.where, n: m.n}) AS mentions "
        "RETURN x {.*} AS sheet, mentions, "
        "[(x)-[r:ROUTED_TO]->(k:Caster) | {caster: k.code, skip: r.skip, doc: left(k.doc, 1500)}][..2] AS old_pipeline_hints",
        key=f"{file}::{sheet}",
    )
    assert rows, f"unknown sheet {file}::{sheet}"
    wb = openpyxl.load_workbook(EXCEL_DIR / file, read_only=True)
    skeleton = []
    for i, row in enumerate(wb[sheet].iter_rows(max_col=16, values_only=True), start=1):
        if any(v is not None for v in row):
            skeleton.append(f"r{i}: " + " | ".join(cell_text(v) for v in row))
        if len(skeleton) == 25:
            break
    wb.close()
    return json.dumps({**rows[0], "skeleton_first_16_cols": skeleton}, ensure_ascii=False, default=str)


@beta_tool
def read_range(file: str, sheet: str, cell_range: str) -> str:
    """Read raw cell values of a rectangular range (max 600 cells). Formulas are returned as text.

    Args:
        file: Workbook file name.
        sheet: Sheet name.
        cell_range: A1-style range, e.g. "A5:L30".
    """
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    assert (max_col - min_col + 1) * (max_row - min_row + 1) <= MAX_CELLS, f"range too large, max {MAX_CELLS} cells"
    wb = openpyxl.load_workbook(EXCEL_DIR / file, read_only=True)
    rows = wb[sheet].iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True)
    out = "\n".join(f"{get_column_letter(min_col)}{r}: " + " | ".join(cell_text(v) for v in row) for r, row in enumerate(rows, start=min_row))
    wb.close()
    return out


def default_anchor(spec: dict) -> dict | None:
    """First text cell of the first label column, from the top of the table: what the layout is recognised by."""
    min_col, min_row, max_col, max_row = range_boundaries(spec["table_range"])
    col = spec["label_columns"][0]
    ci = openpyxl.utils.column_index_from_string(col)
    wb = openpyxl.load_workbook(EXCEL_DIR / spec["file"], read_only=True)
    cells = list(wb[spec["sheet"]].iter_rows(min_row=min_row, max_row=min(max_row, min_row + 10), min_col=ci, max_col=ci, values_only=True))
    wb.close()
    return next(({"cell": f"{col}{min_row + i}", "text": str(v[0]).strip()} for i, v in enumerate(cells)
                 if isinstance(v[0], str) and v[0].strip()), None)


_harvests: dict = {}


@beta_tool
def describe_formulas(file: str, sheet: str = "") -> str:
    """The calculation structure of a workbook recovered from its formulas: per sheet, row rules such as
    "Revenue[row] = Volume[row] * Price[row]" or "SUM([block above])" (with how many rows follow each rule), which sheets feed
    which, and links to workbooks we do not have. Use it on a branch's own spreadsheet to learn its model before mapping
    its rows to concepts, and to ask the branch about inputs that come from outside the file.

    Args:
        file: Workbook file name.
        sheet: Optional sheet name; empty for every sheet (rules covering 3+ rows only).
    """
    path = EXCEL_DIR / file
    key = (str(path), path.stat().st_mtime)
    if key not in _harvests:
        _harvests[key] = formula_graph.harvest(path)
    h = _harvests[key]
    if sheet:
        assert sheet in h["sheets"], f"unknown sheet {sheet}"
        sheets = {sheet: {**h["sheets"][sheet], "rules": h["sheets"][sheet]["rules"][:30]}}
    else:
        sheets = {n: {"formulas": s["formulas"], "rules": [r for r in s["rules"] if r["rows"] >= 3][:8]} for n, s in h["sheets"].items() if s["formulas"]}
    deps = [d for d in h["depends_on"] if not sheet or sheet in (d["sheet"], d["uses"])]
    return json.dumps({"sheets": sheets, "depends_on": deps[:60], "external_links": h["external_links"]}, ensure_ascii=False)


@beta_tool
def propose_extraction_spec(spec_yaml: str) -> str:
    """Propose how to turn a sheet into long rows (one value per row). The spec is executed immediately and checked
    (units, orientation, total rows/columns, region subtotals). Failed checks come back to you: fix the spec and call again.
    Passing specs go to human review.

    Args:
        spec_yaml: YAML with keys file, sheet, table_range (A1, including header rows), header_rows (absolute row numbers),
            label_columns (letters, e.g. [A]), row_dimensions (one name per label column, e.g. [importer]),
            column_dimension (name for the header label, e.g. exporter or period), drop_rows (absolute rows, e.g. unlabeled totals),
            drop_columns (letters, e.g. a Total column), skip_rows_matching (substrings, e.g. [total]),
            skip_label_kinds (e.g. [region] to drop region subtotal rows mixed with countries, else []),
            operators (list of {op, args}: ffill {target: header|labels} for merged/blank-repeated cells;
                subtitle {dimension, rows: {row_number: value}} when one sheet stacks sections, e.g. exports then imports; stack last),
            value_unit (unit from the ingestion referential, e.g. kiloton, kiloton_p2o5, million_ton, usd_per_ton, usd_per_short_ton, percent; aliases like kt, $/t accepted), fixed (dict, e.g. {product: dap, year: 2024}), notes.
            Optional: published (edition date, YYYY-MM-DD: later periods are flagged forecast even without a marker);
            missing_values (provider "no market" markers, e.g. [0, NM, n/a]: use 0 only for prices, never for volumes);
            anchor ({cell, text}; added automatically from the first label cell if you leave it out).
            Every extracted row gets status forecast|actual from the cell formats (e.g. Argus "mmm-yy f").
    """
    spec = yaml.safe_load(spec_yaml)
    if "anchor" not in spec:  # pin the layout so a later edition with shifted rows is re-located or refused
        anchor = default_anchor(spec)
        if anchor:
            spec["anchor"] = anchor
            spec_yaml = spec_yaml.rstrip() + "\n" + yaml.safe_dump({"anchor": anchor}, allow_unicode=True)
    validate(spec)
    rows, ctx = execute(spec)
    results = check(spec, rows, ctx)
    failed = [r for r in results if not r["ok"]]
    status = "rejected" if failed else "proposed"
    folder = SPECS / status
    folder.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S")
    name = f"{Path(spec['file']).stem}__{spec['sheet']}__{stamp}".replace(" ", "_").replace("/", "_")
    (folder / f"{name}.yaml").write_text(spec_yaml, encoding="utf-8")
    (folder / f"{name}.checks.json").write_text(json.dumps({"checks": results, "rows": len(rows), "sample": rows[:20]}, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    report = {"status": status, "rows_extracted": len(rows), "checks": results, "sample_rows": rows[:5]}
    if failed:
        report["next"] = "Fix the failed checks (re-read the sheet with read_range if needed) and call propose_extraction_spec again."
    return json.dumps(report, ensure_ascii=False, default=str)
