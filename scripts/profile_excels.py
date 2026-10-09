"""Profile the structure of market-intel Excels: how "messy" is each sheet? Writes data/excel_profile.json."""
import datetime as dt
import json
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SRC = Path(sys.argv[1])
OUT = Path(__file__).parent.parent / "data" / "excel_profile.json"
SCAN_ROWS = 300


def is_time_label(v) -> bool:
    if isinstance(v, (dt.date, dt.datetime)):
        return True
    if isinstance(v, (int, float)) and 1990 <= v <= 2060 and v == int(v):
        return True
    return isinstance(v, str) and v.strip()[:4].isdigit() and 1990 <= int(v.strip()[:4]) <= 2060 and len(v.strip()) <= 8


def merged_counts(path: Path) -> dict[str, int]:
    """Merged ranges per sheet title, read from the raw XML (read_only openpyxl does not expose them)."""
    with zipfile.ZipFile(path) as z:
        rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))}
        sheets = ET.fromstring(z.read("xl/workbook.xml")).find(f"{{{NS}}}sheets")
        out = {}
        for s in sheets:
            target = rels[s.get(f"{{{REL_NS}}}id")].lstrip("/").removeprefix("xl/")
            out[s.get("name")] = z.read(f"xl/{target}").count(b"<mergeCell ")
        return out


def profile_sheet(ws, merged: int) -> dict:
    rows = [r for _, r in zip(range(SCAN_ROWS), ws.iter_rows(values_only=True))]
    filled = [[v for v in r if v is not None and str(v).strip() != ""] for r in rows]
    header_idx = next((i for i, f in enumerate(filled) if len(f) >= 3 and sum(isinstance(v, str) for v in f) / len(f) >= 0.5), None)
    header_depth = 0
    if header_idx is not None:
        for f in filled[header_idx:]:
            if not f or sum(isinstance(v, str) for v in f) / len(f) < 0.5:
                break
            header_depth += 1
    header = rows[header_idx] if header_idx is not None else ()
    start = header_idx if header_idx is not None else 0
    header_rows = filled[start:start + min(max(header_depth, 1), 3)]
    blank_runs = sum(1 for i in range(1, len(filled)) if not filled[i] and filled[i - 1])
    formulas = 0
    total_rows = 0
    for r in ws.iter_rows(values_only=True):
        row_formulas = sum(1 for v in r if isinstance(v, str) and v.startswith("="))
        formulas += row_formulas
        total_rows += row_formulas > 0 and any(isinstance(v, str) and "total" in v.lower() for v in r[:4])
    strings = lambda cells: sorted({str(v).strip()[:60] for v in cells if isinstance(v, str) and not v.startswith("=") and v.strip()})
    return {
        "sheet": ws.title,
        "hidden": ws.sheet_state != "visible",
        "rows": ws.max_row,
        "cols": ws.max_column,
        "merged_ranges": merged,
        "formulas": formulas,
        "formula_total_rows": total_rows,
        "title": str(next((v for f in filled[:5] for v in f), ""))[:120],
        "header_row": None if header_idx is None else header_idx + 1,
        "header_depth": header_depth,
        "time_in_columns": sum(is_time_label(v) for v in header) >= 5,
        "blocks_in_first_rows": blank_runs + 1,
        "preamble": strings(v for f in filled[:start] for v in f)[:30],
        "header_cells": strings(v for f in header_rows for v in f)[:200],
        "row_labels": strings(v for r in rows[start + len(header_rows):] for v in r[:3])[:300],
    }


def main() -> None:
    result = []
    for path in sorted(SRC.glob("*.xlsx")):
        print("profiling", path.name, flush=True)
        merged = merged_counts(path)
        wb = openpyxl.load_workbook(path, read_only=True)
        sheets = [profile_sheet(ws, merged[ws.title]) for ws in wb.worksheets]
        wb.close()
        result.append({"file": path.name, "size_mb": round(path.stat().st_size / 1e6, 1), "sheets": sheets})
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps(result, indent=1, ensure_ascii=False, default=str), encoding="utf-8")
    print("wrote", OUT)


if __name__ == "__main__":
    main()
